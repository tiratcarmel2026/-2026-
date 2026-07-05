"""Word (.docx) export with proper RTL Hebrew paragraphs."""
from __future__ import annotations

import io

from .export import format_timestamp_short, speaker_display_name


def _set_rtl(paragraph) -> None:
    from docx.oxml import OxmlElement

    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    p_pr.append(bidi)


def build_docx(transcript: dict, speaker_names: dict[str, str], original_filename: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    title = doc.add_heading(f"תמלול: {original_filename}", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl(title)

    for seg in transcript["segments"]:
        name = speaker_display_name(seg["speaker"], speaker_names, transcript["speakers"])
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl(paragraph)
        prefix_run = paragraph.add_run(f"[{format_timestamp_short(seg['start'])}] {name}: ")
        prefix_run.bold = True
        paragraph.add_run(seg["text"])

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
